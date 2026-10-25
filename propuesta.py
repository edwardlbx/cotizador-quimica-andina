"""
Documentos de la PROPUESTA DE PRECIOS (Química Andina S.A.S. — empresa ficticia)
--------------------------------------------------------------------------------
- generar_pdf(datos)      -> bytes del PDF con diseño de propuesta
- excel_erp(datos)      -> bytes de un Excel en el orden de digitación del ERP
- texto_whatsapp(datos)   -> texto corto para pegar en WhatsApp
- guardar_logo_png(ruta)  -> exporta el logo como PNG

`datos` es un dict:
  numero, fecha (date), vigente_hasta (date), asesor, observaciones
  cliente: {nombre, nit, ciudad, condicion_pago_dias}
  lineas:  [{codigo, descripcion, cantidad, um, precio, valor, entrega}, …]
  subtotal
"""
from datetime import date
from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (Flowable, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle, KeepTogether)

# ----------------------------------------------------------------------
# Identidad (ficticia)
# ----------------------------------------------------------------------
EMPRESA = {
    "nombre": "Química Andina S.A.S.",
    "nit": "NIT 901.999.999-0",
    "direccion": "Av. Ejemplo # 00-00, Zona Industrial · Bogotá D.C. (dirección ficticia)",
    "contacto": "Tel. +57 602 000 0000 · ventas@quimicaandina.example",
    "web": "www.quimicaandina.example",
    "lema": "Materias primas para la industria",
}
VERDE = colors.HexColor("#0E6B5C")      # verde andino
VERDE_CLARO = colors.HexColor("#E3F1EC")
AMBAR = colors.HexColor("#D9A23A")      # sol
GRIS = colors.HexColor("#5F5E5A")
GRIS_LINEA = colors.HexColor("#D3D1C7")


# ----------------------------------------------------------------------
# Formatos colombianos
# ----------------------------------------------------------------------
def cop(x) -> str:
    return "$" + f"{float(x):,.0f}".replace(",", ".")

def num(x) -> str:
    s = f"{float(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s.rstrip("0").rstrip(",")

def fecha(d: date) -> str:
    meses = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
    return f"{d.day:02d}-{meses[d.month - 1]}-{d.year}"

def entrega_cliente(texto: str) -> str:
    """Traduce la entrega interna del cotizador a un texto apto para el cliente."""
    t = (texto or "").lower()
    if t.startswith("inmediata"):
        return "Inmediata"
    if t.startswith("pocos días"):
        return "3 a 5 días hábiles"
    if t.startswith("fabricar"):
        dias = "".join(ch for ch in t if ch.isdigit())
        return f"{dias} días hábiles aprox." if dias else "Por confirmar"
    return "Por confirmar"


# ----------------------------------------------------------------------
# Logo: montañas andinas dentro de una gota, con sol
# ----------------------------------------------------------------------
def _dibujar_logo(c, x, y, s):
    """Dibuja el logo con esquina inferior izquierda en (x, y) y tamaño s."""
    c.saveState()
    cx, cy, r = x + s / 2, y + s * 0.42, s * 0.42
    # gota
    p = c.beginPath()
    p.moveTo(cx, y + s)                                   # punta
    p.curveTo(cx + r * 0.35, y + s * 0.80, cx + r, cy + r * 0.55, cx + r, cy)
    p.curveTo(cx + r, cy - r * 0.95, cx - r, cy - r * 0.95, cx - r, cy)
    p.curveTo(cx - r, cy + r * 0.55, cx - r * 0.35, y + s * 0.80, cx, y + s)
    p.close()
    c.setFillColor(VERDE)
    c.drawPath(p, stroke=0, fill=1)
    # sol
    c.setFillColor(AMBAR)
    c.circle(cx + r * 0.38, cy + r * 0.28, r * 0.20, stroke=0, fill=1)
    # montañas
    c.setFillColor(colors.white)
    m = c.beginPath()
    base = cy - r * 0.45
    m.moveTo(cx - r * 0.78, base)
    m.lineTo(cx - r * 0.28, cy + r * 0.22)
    m.lineTo(cx - r * 0.02, cy - r * 0.08)
    m.lineTo(cx + r * 0.22, cy + r * 0.05)
    m.lineTo(cx + r * 0.78, base)
    m.close()
    c.drawPath(m, stroke=0, fill=1)
    # nieve / línea de la cima
    c.setStrokeColor(VERDE)
    c.setLineWidth(max(0.6, s * 0.02))
    c.line(cx - r * 0.42, cy + r * 0.02, cx - r * 0.28, cy + r * 0.10)
    c.restoreState()


class Logo(Flowable):
    def __init__(self, tam=16 * mm):
        super().__init__()
        self.tam = tam
        self.width = self.height = tam

    def draw(self):
        _dibujar_logo(self.canv, 0, 0, self.tam)


def guardar_logo_png(ruta="logo_quimica_andina.png", px=600):
    """Exporta el logo (gota + nombre) como PNG usando Pillow."""
    import math
    from PIL import Image, ImageDraw, ImageFont
    H = px
    def fuente(negrita, tam):
        for n in (("arialbd.ttf", "DejaVuSans-Bold.ttf") if negrita else ("arial.ttf", "DejaVuSans.ttf")):
            try:
                return ImageFont.truetype(n, tam)
            except OSError:
                pass
        return ImageFont.load_default()
    f1, f2 = fuente(True, int(H * 0.30)), fuente(False, int(H * 0.11))
    t1, t2 = "QUÍMICA ANDINA", "MATERIAS PRIMAS PARA LA INDUSTRIA"
    medir = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    W = int(H * 1.05 + max(medir.textlength(t1, font=f1), medir.textlength(t2, font=f2)) + H * 0.1)
    img = Image.new("RGBA", (W, H), (255, 255, 255, 0))
    d = ImageDraw.Draw(img)
    s, ox, oy = H * 0.9, H * 0.05, H * 0.05
    cx, cy, r = ox + s / 2, oy + s * 0.58, s * 0.42
    verde, ambar = (14, 107, 92), (217, 162, 58)
    d.polygon([(cx, oy)] + [(cx + r * math.cos(math.radians(a)), cy + r * math.sin(math.radians(a)))
                            for a in range(-20, 201, 4)], fill=verde)                       # gota
    d.ellipse([cx + r * 0.18, cy - r * 0.48, cx + r * 0.58, cy - r * 0.08], fill=ambar)     # sol
    base = cy + r * 0.45
    d.polygon([(cx - r * 0.78, base), (cx - r * 0.28, cy - r * 0.22), (cx - r * 0.02, cy + r * 0.08),
               (cx + r * 0.22, cy - r * 0.05), (cx + r * 0.78, base)], fill="white")         # montañas
    d.text((H * 1.05, H * 0.20), t1, font=f1, fill=verde)
    d.text((H * 1.07, H * 0.62), t2, font=f2, fill=(95, 94, 90))
    img.save(ruta)
    return ruta


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def _estilos():
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.black)
    return {
        "base": base,
        "peq": ParagraphStyle("peq", parent=base, fontSize=7.5, leading=9.5, textColor=GRIS),
        "empresa": ParagraphStyle("empresa", parent=base, fontName="Helvetica-Bold", fontSize=14,
                                  leading=17, textColor=VERDE),
        "titulo": ParagraphStyle("titulo", parent=base, fontName="Helvetica-Bold", fontSize=13,
                                 leading=16, textColor=colors.white, alignment=TA_RIGHT),
        "subtitulo": ParagraphStyle("subtitulo", parent=base, fontSize=8, leading=10,
                                    textColor=colors.white, alignment=TA_RIGHT),
        "der": ParagraphStyle("der", parent=base, alignment=TA_RIGHT),
        "seccion": ParagraphStyle("seccion", parent=base, fontName="Helvetica-Bold", fontSize=9,
                                  textColor=VERDE, spaceBefore=4, spaceAfter=3),
        "celda": ParagraphStyle("celda", parent=base, fontSize=8, leading=10),
    }


def _pie(c, doc):
    c.saveState()
    w, _ = letter
    c.setStrokeColor(GRIS_LINEA)
    c.line(18 * mm, 16 * mm, w - 18 * mm, 16 * mm)
    c.setFont("Helvetica", 7)
    c.setFillColor(GRIS)
    c.drawString(18 * mm, 12 * mm,
                 "Documento informativo: no constituye oferta comercial en firme. "
                 "La cotización formal se emite desde nuestro sistema ERP.")
    c.drawString(18 * mm, 8.5 * mm, f"{EMPRESA['nombre']} · {EMPRESA['web']} · "
                                    "Empresa ficticia — documento de un proyecto académico.")
    c.drawRightString(w - 18 * mm, 8.5 * mm, f"Página {doc.page}")
    c.restoreState()


def generar_pdf(datos: dict) -> bytes:
    E = _estilos()
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=15 * mm, bottomMargin=22 * mm,
                            title=f"Propuesta de precios {datos['numero']}",
                            author=EMPRESA["nombre"])
    ancho = letter[0] - 36 * mm
    h = []

    # --- Encabezado: logo + empresa | caja de documento ---
    empresa = [Paragraph(EMPRESA["nombre"], E["empresa"]),
               Paragraph(EMPRESA["lema"], E["peq"]),
               Spacer(1, 3),
               Paragraph(f"{EMPRESA['nit']}<br/>{EMPRESA['direccion']}<br/>{EMPRESA['contacto']}", E["peq"])]
    caja = Table([[Paragraph("PROPUESTA DE PRECIOS", E["titulo"])],
                  [Paragraph("Documento informativo", E["subtitulo"])],
                  [Paragraph(f"<b>N.º {datos['numero']}</b>", E["titulo"])]],
                 colWidths=[62 * mm])
    caja.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), VERDE),
                              ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                              ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                              ("TOPPADDING", (0, 0), (-1, 0), 7), ("BOTTOMPADDING", (0, -1), (-1, -1), 7)]))
    enc = Table([[Logo(17 * mm), empresa, caja]], colWidths=[21 * mm, ancho - 21 * mm - 62 * mm, 62 * mm])
    enc.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                             ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    h += [enc, Spacer(1, 6 * mm)]

    # --- Datos del cliente y de la propuesta ---
    cli = datos["cliente"]
    pago = cli.get("condicion_pago_dias")
    pago_txt = "Contado" if not pago else f"Crédito {pago} días"
    filas = [
        ["Cliente", cli["nombre"], "Fecha", fecha(datos["fecha"])],
        ["NIT / C.C.", cli["nit"], "Válida hasta", fecha(datos["vigente_hasta"])],
        ["Ciudad", cli.get("ciudad") or "—", "Asesor", datos.get("asesor") or "—"],
        ["Condición de pago", pago_txt, "Moneda", "Pesos colombianos (COP)"],
    ]
    t = Table([[Paragraph(f"<b>{a}</b>", E["celda"]), Paragraph(str(b), E["celda"]),
                Paragraph(f"<b>{c}</b>", E["celda"]), Paragraph(str(d), E["celda"])] for a, b, c, d in filas],
              colWidths=[30 * mm, ancho / 2 - 30 * mm, 24 * mm, ancho / 2 - 24 * mm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), VERDE_CLARO),
                           ("BOX", (0, 0), (-1, -1), 0.5, GRIS_LINEA),
                           ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    h += [t, Spacer(1, 5 * mm)]

    # --- Productos ---
    enc_tabla = ["#", "Código", "Descripción", "Cantidad", "U.M.", "Precio unit.", "Valor", "Entrega"]
    cuerpo = [[Paragraph(f"<b>{x}</b>", ParagraphStyle("h", parent=E["celda"], textColor=colors.white))
               for x in enc_tabla]]
    for i, l in enumerate(datos["lineas"], 1):
        cuerpo.append([str(i), l["codigo"], Paragraph(l["descripcion"], E["celda"]), num(l["cantidad"]),
                       l.get("um", "KG"), cop(l["precio"]), cop(l["valor"]),
                       Paragraph(entrega_cliente(l.get("entrega")), E["celda"])])
    anchos = [7 * mm, 15 * mm, None, 17 * mm, 11 * mm, 22 * mm, 24 * mm, 26 * mm]
    anchos[2] = ancho - sum(a for a in anchos if a)
    tp = Table(cuerpo, colWidths=anchos, repeatRows=1)
    tp.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), VERDE),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F5F1")]),
        ("ALIGN", (3, 1), (3, -1), "RIGHT"), ("ALIGN", (5, 1), (6, -1), "RIGHT"),
        ("ALIGN", (0, 1), (0, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, -1), (-1, -1), 0.8, VERDE),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    h.append(tp)

    # --- Total ---
    tot = Table([["", "Subtotal (sin IVA)", cop(datos["subtotal"])]],
                colWidths=[ancho - 74 * mm, 40 * mm, 34 * mm])
    tot.setStyle(TableStyle([("FONTNAME", (1, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (1, 0), (-1, 0), 10),
                             ("TEXTCOLOR", (1, 0), (-1, 0), VERDE), ("ALIGN", (1, 0), (-1, 0), "RIGHT"),
                             ("LINEBELOW", (1, 0), (-1, 0), 1.2, AMBAR), ("TOPPADDING", (0, 0), (-1, -1), 6)]))
    h += [tot, Spacer(1, 6 * mm)]

    # --- Condiciones ---
    cond = [
        f"Propuesta válida hasta el {fecha(datos['vigente_hasta'])}. Precios en pesos colombianos; no incluyen IVA.",
        "Precios y disponibilidad sujetos a confirmación al momento de formalizar la cotización en nuestro ERP.",
        "Los tiempos de entrega son estimados y se cuentan en días hábiles desde la confirmación del pedido.",
        "Condición de pago según lo pactado con el cliente.",
    ]
    if datos.get("observaciones"):
        cond.append(f"Observaciones: {datos['observaciones']}")
    bloque = [Paragraph("Condiciones", E["seccion"])] + [Paragraph("• " + x, E["base"]) for x in cond]
    bloque += [Spacer(1, 8 * mm),
               Paragraph(f"<b>{datos.get('asesor') or ''}</b><br/>Asesor comercial · {EMPRESA['contacto']}",
                         E["peq"])]
    h.append(KeepTogether(bloque))

    doc.build(h, onFirstPage=_pie, onLaterPages=_pie)
    return buf.getvalue()


# ----------------------------------------------------------------------
# Excel para digitar en el ERP
# ----------------------------------------------------------------------
def excel_erp(datos: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Para ERP"
    cab = [("Propuesta", datos["numero"]), ("Fecha", datos["fecha"].strftime("%d/%m/%Y")),
           ("NIT cliente", datos["cliente"]["nit"]), ("Cliente", datos["cliente"]["nombre"]),
           ("Asesor", datos.get("asesor") or "")]
    for i, (k, v) in enumerate(cab, 1):
        ws.cell(i, 1, k).font = Font(name="Arial", bold=True)
        ws.cell(i, 2, v).font = Font(name="Arial")
    fila = len(cab) + 2
    titulos = ["Referencia", "Descripción", "Cantidad", "Precio unitario", "Valor"]
    for j, t in enumerate(titulos, 1):
        c = ws.cell(fila, j, t)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0E6B5C")
        c.alignment = Alignment(horizontal="center")
    for l in datos["lineas"]:
        fila += 1
        ws.cell(fila, 1, l["codigo"])
        ws.cell(fila, 2, l["descripcion"])
        ws.cell(fila, 3, float(l["cantidad"])).number_format = "#,##0.00"
        ws.cell(fila, 4, float(l["precio"])).number_format = "$#,##0"
        ws.cell(fila, 5, f"=C{fila}*D{fila}").number_format = "$#,##0"
    primera = len(cab) + 3
    ws.cell(fila + 1, 4, "Subtotal").font = Font(name="Arial", bold=True)
    tot = ws.cell(fila + 1, 5, f"=SUM(E{primera}:E{fila})")
    tot.font = Font(name="Arial", bold=True)
    tot.number_format = "$#,##0"
    for col, w in zip("ABCDE", [14, 50, 12, 16, 16]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=primera, max_row=fila):
        for c in row:
            c.font = Font(name="Arial")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------
# Texto para WhatsApp
# ----------------------------------------------------------------------
def texto_whatsapp(datos: dict) -> str:
    lineas = [f"*{EMPRESA['nombre']}* – Propuesta de precios *{datos['numero']}*",
              f"Cliente: {datos['cliente']['nombre']}", ""]
    for l in datos["lineas"]:
        lineas.append(f"• {l['descripcion']} – {num(l['cantidad'])} {l.get('um', 'kg').lower()} – "
                      f"{cop(l['precio'])}/{l.get('um', 'kg').lower()} – entrega: {entrega_cliente(l.get('entrega')).lower()}")
    lineas += ["", f"*Subtotal: {cop(datos['subtotal'])} + IVA*",
               f"Válida hasta el {fecha(datos['vigente_hasta'])}.",
               "_Precios informativos; la cotización formal se envía desde nuestro ERP._"]
    return "\n".join(lineas)


if __name__ == "__main__":                       # genera un ejemplo y el logo
    ejemplo = {
        "numero": "PP-2026-0001", "fecha": date.today(), "vigente_hasta": date.today(),
        "asesor": "Camila Andrea Morales Ruiz", "observaciones": None,
        "cliente": {"nombre": "CLIENTE-0001", "nit": "900.000.001", "ciudad": "Cali",
                    "condicion_pago_dias": 30},
        "lineas": [{"codigo": "0917", "descripcion": "PRODUCTO-0917",
                    "cantidad": 10, "um": "KG", "precio": 171000, "valor": 1710000,
                    "entrega": "Inmediata (almacén)"}],
        "subtotal": 1710000,
    }
    open("propuesta_ejemplo.pdf", "wb").write(generar_pdf(ejemplo))
    guardar_logo_png()
    print("propuesta_ejemplo.pdf y logo_quimica_andina.png generados")
